"""
Data loading + UI helpers for US Pizza Student Voucher Dashboard
"""

from pathlib import Path
import re
import pandas as pd
import streamlit as st
import plotly.express as px
from openpyxl import load_workbook
from styles import apply_theme, render_top_navigation, render_footer

# Excel sits next to this file
WORKBOOK = Path(__file__).parent / "Student_Voucher_Route_Plan_2026_v3.xlsx"


def google_sheets_enabled() -> bool:
    """Use Google Sheets when its ID and service-account credentials are configured."""
    return bool(st.secrets.get("GOOGLE_SHEET_ID") and st.secrets.get("gcp_service_account"))


@st.cache_resource(show_spinner=False)
def _sheets_client():
    import gspread
    return gspread.service_account_from_dict(dict(st.secrets["gcp_service_account"]))


def _spreadsheet():
    return _sheets_client().open_by_key(str(st.secrets["GOOGLE_SHEET_ID"]))


def _values_frame(values) -> pd.DataFrame:
    """Make a rectangular frame from Google Sheets values (which omit trailing cells)."""
    width = max((len(row) for row in values), default=0)
    return pd.DataFrame([list(row) + [None] * (width - len(row)) for row in values])


def clean(value) -> str:
    return "" if pd.isna(value) else re.sub(r"\s+", " ", str(value)).strip()


def outlet_key(value) -> str:
    return re.sub(r"[^a-z0-9]", "", clean(value).lower())


def region_for_state(state) -> str:
    state = clean(state).lower()
    if state in {"selangor", "kuala lumpur", "putrajaya"}:
        return "KV"
    if state in {"perlis", "kedah", "pulau pinang", "penang", "perak"}:
        return "Northern"
    if state in {"negeri sembilan", "melaka", "johor"}:
        return "Southern"
    return "Unassigned"


@st.cache_data(show_spinner=False)
def load_outlet_roster(path_text: str) -> pd.DataFrame:
    """Load outlet-to-state values used only to infer each school's region."""
    if google_sheets_enabled():
        values = _spreadsheet().worksheet("Manpower & Outlets").get_all_values()
        raw = _values_frame(values)
        if len(raw) < 3:
            raise ValueError("The Manpower & Outlets sheet needs its column headers on row 3.")
        raw.columns = [clean(value).upper() or f"COLUMN_{i + 1}" for i, value in enumerate(raw.iloc[2].tolist())]
        roster = raw.iloc[3:].copy()
    else:
        path = Path(path_text)
        roster = pd.read_excel(path, sheet_name="Manpower & Outlets", header=2)
    required = {"OUTLET", "STATE"}
    if not required.issubset(roster.columns):
        raise ValueError("The Manpower & Outlets tab must contain OUTLET and STATE.")
    roster = roster[list(required)].copy()
    roster["OUTLET"] = roster["OUTLET"].map(clean)
    roster["STATE"] = roster["STATE"].map(clean)
    roster = roster[roster["OUTLET"] != ""].copy()
    roster["OUTLET_KEY"] = roster["OUTLET"].map(outlet_key)
    roster["REGION"] = roster["STATE"].map(region_for_state)
    return roster


@st.cache_data(show_spinner=False)
def load_routes(path_text: str) -> pd.DataFrame:
    path = Path(path_text)
    if not google_sheets_enabled() and not path.exists():
        raise FileNotFoundError(
            "Google Sheets mode is not active and no local workbook is available. "
            "In Streamlit Community Cloud, check Settings → Secrets for both "
            "GOOGLE_SHEET_ID and the [gcp_service_account] section."
        )

    frames = []
    if google_sheets_enabled():
        source_sheets = [(ws.title, _values_frame(ws.get_all_values())) for ws in _spreadsheet().worksheets()]
    else:
        book = pd.ExcelFile(path, engine="openpyxl")
        source_sheets = [(name, book.parse(sheet_name=name, header=None)) for name in book.sheet_names]
    outlet_roster = load_outlet_roster(str(path))
    outlet_lookup = outlet_roster.set_index("OUTLET_KEY").to_dict("index")

    for sheet, raw in source_sheets:
        header_row = None
        for i in range(min(12, len(raw))):
            labels = {clean(x).upper() for x in raw.iloc[i].tolist()}
            if {"STOP", "SUGGESTED DAY", "LEVEL", "NAMA SEKOLAH", "TOTAL"}.issubset(labels):
                header_row = i
                break
        if header_row is None:
            continue

        headers = [clean(c).upper() or f"COLUMN_{j + 1}" for j, c in enumerate(raw.iloc[header_row].tolist())]
        df = raw.iloc[header_row + 1:].copy()
        df.columns = headers
        if not {"LEVEL", "NAMA SEKOLAH", "TOTAL"}.issubset(df.columns):
            continue

        df = df[df["NAMA SEKOLAH"].notna()].copy()
        df["MANAGER"] = sheet.strip()
        if "NEAREST OUTLET" not in df:
            df["NEAREST OUTLET"] = ""
        if "AREA" not in df:
            df["AREA"] = ""
        df["AREA"] = df["AREA"].map(clean)
        df.loc[df["AREA"] == "", "AREA"] = df.loc[df["AREA"] == "", "NEAREST OUTLET"].map(clean)
        df["AREA_KEY"] = df["AREA"].map(outlet_key)
        df["OUTLET_KEY"] = df["NEAREST OUTLET"].map(outlet_key)
        df["OUTLET_STATE"] = df["OUTLET_KEY"].map(
            lambda key: outlet_lookup.get(key, {}).get("STATE", "")
        )
        df["REGION"] = df["OUTLET_STATE"].map(region_for_state)
        df["_SHEET"] = sheet
        df["_HEADER_ROW"] = header_row + 1
        # `raw` keeps its original worksheet row index, so index + 1 is the Excel row.
        df["_EXCEL_ROW"] = df.index + 1
        df["LEVEL"] = df["LEVEL"].map(clean).str.title()

        for col in ("TOTAL", "QTY GIVEN"):
            if col not in df:
                df[col] = 0
            df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

        for col in ("STATUS", "STAFF NAME", "DATE", "CONTACT NUMBER", "EMAIL ADDRESS", "REJECTION REASON", "PERSON MET", "EVENT REMARKS", "QTY NOT PROVIDED"):
            if col not in df:
                df[col] = ""
            df[col] = df[col].map(clean)

        df["STATUS_KEY"] = df["STATUS"].str.lower().str.replace(r"[^a-z]", "", regex=True)
        df["REJECTED"] = df["STATUS_KEY"].isin({"reject", "rejected", "declined", "unsuccessful"})
        df["SUCCESS"] = df["STATUS_KEY"].isin({
            "received", "success", "successful", "completed", "complete", "delivered", "done"
        })
        df["GIVEN"] = (df["QTY GIVEN"] > 0) | df["SUCCESS"]
        frames.append(df)

    if not frames:
        raise ValueError("No route tables found. Expected columns: LEVEL, NAMA SEKOLAH, TOTAL.")
    return pd.concat(frames, ignore_index=True)


def page_setup(title: str) -> pd.DataFrame:
    st.set_page_config(
        page_title=f"US Pizza · {title}",
        page_icon="🍕",
        layout="wide",
        initial_sidebar_state="collapsed",
    )
    apply_theme()
    render_top_navigation()

    try:
        sheets_enabled = google_sheets_enabled()
    except Exception as exc:
        st.error(f"Could not read Streamlit secrets: {exc}")
        st.stop()

    with st.sidebar:
        if st.button("🔄 Refresh data", use_container_width=True):
            load_routes.clear()
            load_outlet_roster.clear()
            st.rerun()
        st.markdown("---")
        st.caption("Updates save to Google Sheets." if sheets_enabled else "Local updates save to the Excel workbook.")

    try:
        routes = load_routes(str(WORKBOOK))
    except Exception as exc:
        # Include the exception type and repr because some Google API exceptions
        # have an empty str(exc), which otherwise leaves the user with a blank error.
        st.error(f"Could not load dashboard data ({type(exc).__name__}): {exc!r}")
        if sheets_enabled:
            st.info("Google Sheets settings were detected. Check that the service account has Editor access and that the required worksheet tabs still exist.")
        else:
            st.info("The app is not detecting its Google Sheets settings. Check this app’s Streamlit Cloud Settings → Secrets, then save and restart the app.")
        st.stop()

    return routes


def metric_cards(df: pd.DataFrame, include_vouchers: bool = True):
    total_schools = len(df)
    primary = int((df["LEVEL"].str.lower() == "primary").sum())
    secondary = int((df["LEVEL"].str.lower() == "secondary").sum())
    given_schools = int(df["GIVEN"].sum())
    rejected = int(df["REJECTED"].sum())
    left = max(total_schools - given_schools, 0)
    if not include_vouchers:
        received = int(df["SUCCESS"].sum())
        visited = received + rejected
        schools_left = max(total_schools - visited, 0)
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Total schools", f"{total_schools:,}")
        c2.metric("Primary schools", f"{primary:,}")
        c3.metric("Secondary schools", f"{secondary:,}")
        c4.metric("Schools left to visit", f"{schools_left:,}")
        c1, c2, c3 = st.columns(3)
        c1.metric("Received schools", f"{received:,}")
        c2.metric("Rejected schools", f"{rejected:,}")
        c3.metric("Visited (received + rejected)", f"{visited:,}")
        return

    c1, c2, c3 = st.columns(3)
    c1.metric("Total schools", f"{total_schools:,}")
    c2.metric("Primary schools", f"{primary:,}")
    c3.metric("Secondary schools", f"{secondary:,}")

    c1, c2, c3, c4 = st.columns(4)
    if include_vouchers:
        c1.metric("Received schools", f"{given_schools:,}")
        c2.metric("Rejected schools", f"{rejected:,}")
        c3.metric("Schools left", f"{left:,}")
        c4.metric("% schools received", f"{(given_schools/total_schools*100) if total_schools else 0:.1f}%")
    if include_vouchers:
        planned = int(df["TOTAL"].sum())
        issued = int(df["QTY GIVEN"].sum())
        balance = max(planned - issued, 0)
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Planned vouchers*", f"{planned:,}")
        c2.metric("Vouchers given", f"{issued:,}")
        c3.metric("Voucher balance*", f"{balance:,}")
        pct = issued / planned if planned else 0
        c4.metric("% vouchers given", f"{pct*100:.1f}%")


def level_table(df: pd.DataFrame, include_vouchers: bool = True):
    rows = []
    for level in ["Primary", "Secondary"]:
        part = df[df["LEVEL"].str.lower() == level.lower()]
        if part.empty:
            continue
        received = int(part["SUCCESS"].sum())
        rejected = int(part["REJECTED"].sum())
        visited = received + rejected
        if include_vouchers:
            summary = {
                "Level": level,
                "Schools": len(part),
                "Received schools": int(part["GIVEN"].sum()),
                "Schools left": max(len(part) - int(part["GIVEN"].sum()), 0),
                "Rejected schools": rejected,
            }
        else:
            summary = {
                "Level": level,
                "Schools": len(part),
                "Received": received,
                "Rejected": rejected,
                "Visited (received + rejected)": visited,
                "Schools left": max(len(part) - visited, 0),
            }
        if include_vouchers:
            planned = int(part["TOTAL"].sum())
            issued = int(part["QTY GIVEN"].sum())
            summary.update({
                "Planned vouchers*": planned,
                "Vouchers given": issued,
                "Voucher balance*": max(planned - issued, 0),
            })
        rows.append(summary)
    if rows:
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)


def manager_table(df: pd.DataFrame, include_vouchers: bool = True):
    rows = []
    for manager, part in df.groupby("MANAGER", dropna=False):
        received = int(part["SUCCESS"].sum())
        rejected = int(part["REJECTED"].sum())
        visited = received + rejected
        summary = {
            "Route manager": manager,
            "Region": ", ".join(sorted(part["REGION"].dropna().unique())),
            "Primary schools": int((part["LEVEL"].str.lower() == "primary").sum()),
            "Secondary schools": int((part["LEVEL"].str.lower() == "secondary").sum()),
            "Received schools": int(part["GIVEN"].sum()) if include_vouchers else received,
            "Schools left": max(len(part) - (int(part["GIVEN"].sum()) if include_vouchers else visited), 0),
            "Rejected schools": rejected,
        }
        if not include_vouchers:
            summary["Visited (received + rejected)"] = visited
        if include_vouchers:
            planned = int(part["TOTAL"].sum())
            issued = int(part["QTY GIVEN"].sum())
            summary.update({
                "Planned vouchers*": planned,
                "Vouchers given": issued,
                "Voucher balance*": max(planned - issued, 0),
            })
        rows.append(summary)
    st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)


def overview_filters(df: pd.DataFrame) -> pd.DataFrame:
    status_col, region_col, pic_col = st.columns([1, 1, 1.5])
    status = status_col.selectbox("School result", ["All", "Received", "Rejected", "Pending", "Visited"], key="overview_status")
    regions = ["All", "KV", "Northern", "Southern"]
    region = region_col.selectbox("Region", regions, key="overview_region")
    manager_source = df if region == "All" else df[df["REGION"] == region]
    managers = ["All"] + sorted(manager_source["MANAGER"].dropna().unique().tolist())
    manager = pic_col.selectbox("Manager", managers, key="overview_manager")

    result = df.copy()
    if region != "All":
        result = result[result["REGION"] == region]
    if manager != "All":
        result = result[result["MANAGER"] == manager]
    st.caption(f"Showing {len(result):,} routed schools · {region} · {manager}")
    st.caption(f"School result filter applies to the school list: {status}.")
    return result, (region, manager), status


def route_filters(df: pd.DataFrame, key: str, include_manager: bool = True, show_count: bool = True) -> pd.DataFrame:
    """Filter a school-level or manager page by region and, when needed, manager."""
    if include_manager:
        region_col, manager_col = st.columns([1, 1.4])
    else:
        region_col = st.columns(1)[0]
        manager_col = None

    region_options = ["All"] + sorted(df["REGION"].dropna().unique().tolist())
    region = region_col.selectbox("Region", region_options, key=f"{key}_region")
    result = df.copy() if region == "All" else df[df["REGION"] == region].copy()

    manager = "All"
    if include_manager:
        manager_options = ["All"] + sorted(result["MANAGER"].dropna().unique().tolist())
        manager = manager_col.selectbox("Manager", manager_options, key=f"{key}_manager")
        if manager != "All":
            result = result[result["MANAGER"] == manager]

    if show_count:
        st.caption(f"Showing {len(result):,} schools · {region} · {manager}")
    return result


def region_table(df: pd.DataFrame):
    rows = []
    for region in ["KV", "Northern", "Southern", "Unassigned"]:
        part = df[df["REGION"] == region]
        if part.empty:
            continue
        rows.append({
            "Region": region,
            "Total schools": len(part),
            "Primary": int((part["LEVEL"].str.lower() == "primary").sum()),
            "Secondary": int((part["LEVEL"].str.lower() == "secondary").sum()),
            "Received": int(part["SUCCESS"].sum()),
            "Rejected": int(part["REJECTED"].sum()),
            "Visited": int(part["SUCCESS"].sum() + part["REJECTED"].sum()),
            "Schools left": max(len(part) - int(part["SUCCESS"].sum() + part["REJECTED"].sum()), 0),
        })
    if rows:
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)


def person_progress_table(routes: pd.DataFrame):
    """Summarize workload using the manager name on each route sheet."""
    rows = []
    for person, part in routes.groupby("MANAGER", dropna=False):
        received = int(part["SUCCESS"].sum())
        rejected = int(part["REJECTED"].sum())
        visited = received + rejected
        rows.append({
            "Manager": clean(person) or "Unassigned",
            "Areas": ", ".join(sorted(part["AREA"].dropna().map(clean).unique())),
            "Total schools": len(part),
            "Primary": int((part["LEVEL"].str.lower() == "primary").sum()),
            "Secondary": int((part["LEVEL"].str.lower() == "secondary").sum()),
            "Received": received,
            "Rejected": rejected,
            "Visited": visited,
            "Schools left": max(len(part) - visited, 0),
        })
    if rows:
        summary = pd.DataFrame(rows).sort_values("Manager")
        wrapped_summary = summary.style.set_properties(
            subset=["Areas"],
            **{
                "white-space": "normal",
                "overflow-wrap": "anywhere",
                "min-width": "260px",
                "max-width": "420px",
            },
        )
        st.table(wrapped_summary)


def area_progress_table(routes: pd.DataFrame, filters):
    region, manager = filters
    scoped_routes = routes.copy()
    if region != "All":
        scoped_routes = scoped_routes[scoped_routes["REGION"] == region]
    if manager != "All":
        scoped_routes = scoped_routes[scoped_routes["MANAGER"] == manager]
    rows = []
    area_groups = scoped_routes.groupby(["AREA", "REGION"], dropna=False)
    for (area_name, area_region), part in area_groups:
        total = len(part)
        received = int(part["SUCCESS"].sum())
        rejected = int(part["REJECTED"].sum())
        visited = received + rejected
        rows.append({
            "Area": area_name,
            "Region": area_region,
            "Manager": ", ".join(sorted(part["MANAGER"].dropna().unique())),
            "Route schools": total,
            "Received": received,
            "Rejected": rejected,
            "Visited (received + rejected)": visited,
            "Schools left": max(total - visited, 0),
        })

    st.caption("Visited counts schools marked Received or Rejected. Area names come from the route sheets.")
    st.dataframe(pd.DataFrame(rows).sort_values(["Region", "Area"]), use_container_width=True, hide_index=True, height=340)


def progress_pie(df: pd.DataFrame, title: str, mode: str):
    """Create a labeled donut chart for school results or voucher balance."""
    if mode == "schools":
        received = int(df["SUCCESS"].sum())
        rejected = int(df["REJECTED"].sum())
        pending = max(len(df) - received - rejected, 0)
        labels = ["Received schools", "Rejected schools", "Pending schools"]
        values = [received, rejected, pending]
        colors = {"Received schools": "#C8202F", "Rejected schools": "#111111", "Pending schools": "#D9D9D9"}
    else:
        planned = int(df["TOTAL"].sum())
        issued = int(df["QTY GIVEN"].sum())
        labels = ["Vouchers given", "Voucher balance"]
        values = [min(issued, planned), max(planned - issued, 0)]
        colors = {"Vouchers given": "#C8202F", "Voucher balance": "#D9D9D9"}
    data = pd.DataFrame({"Category": labels, "Count": values})
    fig = px.pie(data, values="Count", names="Category", hole=0.52,
                 color="Category", color_discrete_map=colors)
    fig.update_traces(textposition="inside", textinfo="label+percent",
                      hovertemplate="%{label}: %{value:,}<extra></extra>")
    fig.update_layout(template="plotly_white", title=title,
                      title_font=dict(color="#111111", size=17),
                      height=260,
                      margin=dict(t=48, b=10, l=8, r=8),
                      legend=dict(orientation="h", yanchor="bottom", y=-0.15,
                                  xanchor="center", x=0.5),
                      paper_bgcolor="rgba(0,0,0,0)", font=dict(color="#111111"))
    return fig


def filtered_school_list(
    df: pd.DataFrame,
    key: str,
    show_status_filter: bool = True,
    include_distribution: bool = True,
    status_filter: str | None = None,
):
    if show_status_filter:
        left, right = st.columns(2)
        status = left.selectbox(
            "Filter by result",
            ["All results", "Received", "Rejected", "Pending", "Visited"],
            key=f"status_{key}",
        )
        search = right.text_input("Search school", key=f"search_{key}").strip().lower()
    elif status_filter is not None:
        status = status_filter
        search = st.text_input("Search school", key=f"search_{key}").strip().lower()
    else:
        status = "All results"
        search = st.text_input("Search school", key=f"search_{key}").strip().lower()

    result = df.copy()
    if status == "Received":
        result = result[result["SUCCESS"] & ~result["REJECTED"]]
    elif status == "Rejected":
        result = result[result["REJECTED"]]
    elif status == "Pending":
        result = result[~result["SUCCESS"] & ~result["REJECTED"]]
    elif status == "Visited":
        result = result[result["SUCCESS"] | result["REJECTED"]]

    if search:
        result = result[result["NAMA SEKOLAH"].astype(str).str.lower().str.contains(search, na=False)]

    cols = [
        "REGION", "MANAGER",
        "LEVEL", "NAMA SEKOLAH", "AREA", "STATUS",
    ]
    if include_distribution:
        cols += [
            "QTY GIVEN", "PERSON MET", "STAFF NAME", "CONTACT NUMBER", "EMAIL ADDRESS",
            "REJECTION REASON", "EVENT REMARKS",
        ]
    cols = [column for column in cols if column in result.columns]
    st.dataframe(result[cols].rename(columns={"AREA": "Area"}), use_container_width=True, hide_index=True)


def save_school_update(
    row: pd.Series, status: str, qty: int, pic: str, phone: str, email: str,
    reason: str, person_met: str = "", event_remarks: str = "",
    qty_not_provided: bool = False,
):
    """Write one submitted school update into its original manager route row."""
    from datetime import datetime
    from zoneinfo import ZoneInfo

    date_value = datetime.now(ZoneInfo("Asia/Kuala_Lumpur")).strftime("%Y-%m-%d") if status != "Pending" else ""
    values = {
        "QTY GIVEN": qty,
        "STAFF NAME": pic,
        "CONTACT NUMBER": phone,
        "EMAIL ADDRESS": email,
        "STATUS": status,
        "REJECTION REASON": reason,
        "DATE": date_value,
        "PERSON MET": person_met,
        "EVENT REMARKS": event_remarks,
        "QTY NOT PROVIDED": "Yes" if qty_not_provided else "",
    }

    if google_sheets_enabled():
        from gspread.utils import rowcol_to_a1

        ws = _spreadsheet().worksheet(str(row["_SHEET"]))
        header = int(row["_HEADER_ROW"])
        current_headers = ws.row_values(header)
        headers = {clean(value).upper(): index + 1 for index, value in enumerate(current_headers) if clean(value)}
        missing_headers = [name for name in values if name not in headers]
        required_columns = len(current_headers) + len(missing_headers)
        if required_columns > ws.col_count:
            # Imported workbook tabs can be limited to A:Z. Expand the grid
            # before writing new tracking fields such as QTY NOT PROVIDED.
            ws.add_cols(required_columns - ws.col_count)
        updates = []
        for name, value in values.items():
            column = headers.get(name)
            if column is None:
                column = len(current_headers) + 1
                updates.append({"range": rowcol_to_a1(header, column), "values": [[name]]})
                current_headers.append(name)
                headers[name] = column
            updates.append({"range": rowcol_to_a1(int(row["_EXCEL_ROW"]), column), "values": [[value]]})
        ws.batch_update(updates, value_input_option="USER_ENTERED")
        load_routes.clear()
        load_outlet_roster.clear()
        return

    path = WORKBOOK
    wb = load_workbook(path)
    ws = wb[row["_SHEET"]]
    header = int(row["_HEADER_ROW"])
    headers = {
        clean(ws.cell(header, c).value).upper(): c
        for c in range(1, ws.max_column + 1)
        if ws.cell(header, c).value is not None
    }
    for name, value in values.items():
        col = headers.get(name)
        if col is None:
            col = ws.max_column + 1
            ws.cell(header, col).value = name
            headers[name] = col
        ws.cell(int(row["_EXCEL_ROW"]), col).value = value

    # Save to a temporary sibling file, then replace the workbook after a complete save.
    temp_path = path.with_name(path.stem + ".dashboard-tmp" + path.suffix)
    try:
        wb.save(temp_path)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    load_routes.clear()

